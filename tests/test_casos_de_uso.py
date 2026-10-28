"""Casos de uso ponta a ponta com SQLite real numa pasta temporária; e-mail e geração de folhas simulados."""
import os
import shutil
import tempfile
import unittest

import openpyxl
import pymupdf
from PIL import Image

from folha_ponto.dominio.entidades import ConfiguracaoEmail, MetodoIdentificacao, Servidor, StatusEnvio, TipoOperacao
from folha_ponto.dominio.regras import ErroDominio, MatriculaDuplicada, PlanilhaInvalida
from folha_ponto.infraestrutura.container import montar_servicos
from folha_ponto.infraestrutura.leitor_folhas import LeitorFolhasTesseract, inclinacao

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class EmailFalso:
    def __init__(self):
        self.enviados = []

    def abrir(self, config):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        pass

    def enviar(self, destino, assunto, corpo, anexo=None):
        if "recusa" in destino:
            return False, "Endereço recusado"
        self.enviados.append((destino, os.path.basename(anexo)))
        return True, "ok"

    def testar(self, config):
        return True, "ok"


class GeradorFalso:
    """Gera PDFs com texto (nome, matrícula, referência) sem precisar do Word."""

    def __init__(self, pasta):
        self.pasta = pasta

    def gerar_lote(self, servidores, competencia, progresso, datas_especiais=None):
        arquivos = {}
        for n, s in enumerate(servidores):
            progresso(n, len(servidores), s.nome)
            destino = os.path.join(self.pasta, f"Folha_{s.matricula}.pdf")
            doc = pymupdf.open()
            doc.new_page().insert_text((50, 72), f"NOME DO SERVIDOR: {s.nome}\nMATRICULA: {s.matricula}\n"
                                                 f"REFERENCIA: {competencia}")
            doc.save(destino)
            doc.close()
            arquivos[s.id] = (destino.replace(".pdf", ".docx"), destino)
        return arquivos, []


def criar_planilha(caminho, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome", "Matrícula", "CPF", "E-mail", "Carga Horária", "Acumula cargo (Sim/Não)", "Cargo"])
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)
    return caminho


class TestCasosDeUso(unittest.TestCase):
    COMP = "FEVEREIRO/2026"

    def setUp(self):
        self.pasta = tempfile.mkdtemp()
        os.makedirs(os.path.join(self.pasta, "recursos"), exist_ok=True)
        shutil.copy(os.path.join(RAIZ, "recursos", "Modelo de folha de ponto - Exemplo.docx"),
                    os.path.join(self.pasta, "recursos"))
        self.email = EmailFalso()
        self.s = montar_servicos(self.pasta, servico_email=self.email, gerador_folhas=GeradorFalso(self.pasta))
        self.s.configuracoes.salvar_email(ConfiguracaoEmail(usuario="digep@gmail.com", senha="x"))
        self.planilha = criar_planilha(os.path.join(self.pasta, "servidores.xlsx"), [
            ["Ana Beatriz Oliveira", "17289104", "444.555.666-04", "ana@teste.br", "40h", "Sim", "Técnico"],
            ["Carlos Eduardo Santos", 17289105, "", "carlos@recusa.br", "20h", "Não", None],
            ["Maria Silva", "17289102", "", "", "40h", "Sim", None],
            [None, "999", None, None, None, None, None],   # sem nome: ignorada
            [None, None, None, None, None, None, None],    # vazia
        ])

    def tearDown(self):
        shutil.rmtree(self.pasta, ignore_errors=True)

    def test_importacao_e_planilha_base(self):
        r = self.s.servidores.definir_planilha_base(self.planilha)
        self.assertEqual((r.novos, r.atualizados, len(r.ids)), (3, 0, 3))
        self.assertEqual(r.ignoradas, ["linha 5 (sem nome)"])
        carlos = self.s.servidores.listar("carlos")[0]
        self.assertEqual(carlos.matricula, "17289105")          # número do Excel sem ".0"
        self.assertEqual(self.s.servidores.listar("tecnico"), [])  # busca não olha o cargo
        self.assertEqual(self.s.servidores.planilha_base().quantidade, 3)
        repetida = self.s.servidores.importar_planilha(self.planilha)  # nada mudou
        self.assertEqual((repetida.novos, repetida.atualizados, len(repetida.ids)), (0, 0, 3))

    def test_planilha_sem_colunas_obrigatorias(self):
        caminho = os.path.join(self.pasta, "ruim.xlsx")
        wb = openpyxl.Workbook()
        wb.active.append(["Qualquer", "Coisa"])
        wb.save(caminho)
        with self.assertRaises(PlanilhaInvalida):
            self.s.servidores.importar_planilha(caminho)

    def test_resumo_previo_antes_de_gravar(self):
        self.s.servidores.importar_planilha(self.planilha)
        nova = criar_planilha(os.path.join(self.pasta, "nova.xlsx"), [
            ["Ana Beatriz Oliveira", "17289104", "444.555.666-04", "ana@teste.br", "40h", "Sim", "Técnico"],   # igual
            ["Carlos Eduardo Santos", "17289105", "", "carlos.novo@teste.br", "20h", "Não", None],            # mudou
            ["Daniela Rocha", "17289200", "", "dani@teste.br", "40h", "Não", None],                           # nova
            ["Daniela Rocha Repetida", "17289200", "", "x@teste.br", "40h", "Não", None],                     # duplicada
            [None, "999", None, None, None, None, None],                                                      # sem nome
        ])
        resumo = self.s.servidores.analisar_planilha(nova)
        self.assertEqual((resumo.novos, resumo.atualizados, resumo.sem_alteracao, len(resumo.ignoradas)), (1, 1, 1, 2))
        self.assertEqual(self.s.servidores.contar(), 3)  # nada gravado ainda
        resultado = self.s.servidores.confirmar_importacao(resumo)
        self.assertEqual((resultado.novos, resultado.atualizados), (1, 1))
        self.assertEqual(self.s.servidores.contar(), 4)
        self.assertEqual(self.s.servidores.listar("daniela")[0].nome, "Daniela Rocha")

    def test_validacao_de_formato_e_cabecalho(self):
        txt = os.path.join(self.pasta, "servidores.csv")
        with open(txt, "w") as f:
            f.write("Nome;Matrícula")
        with self.assertRaises(PlanilhaInvalida) as e:
            self.s.servidores.analisar_planilha(txt)
        self.assertIn(".xlsx", str(e.exception))
        incompleta = os.path.join(self.pasta, "incompleta.xlsx")
        wb = openpyxl.Workbook()
        wb.active.append(["Nome", "Matrícula", "E-mail"])
        wb.active.append(["Ana", "1", "a@b.c"])
        wb.save(incompleta)
        with self.assertRaises(PlanilhaInvalida) as e:
            self.s.servidores.analisar_planilha(incompleta)
        self.assertIn("CPF", str(e.exception))
        self.assertIn("Carga Horária", str(e.exception))

    def test_inclusao_e_remocao_manual(self):
        self.s.servidores.importar_planilha(self.planilha)
        novo = Servidor(nome="Zélia Manual", matricula="77770", cpf="111.222.333-44", email="z@teste.br")
        self.s.servidores.salvar(novo)
        self.assertEqual([s.nome for s in self.s.servidores.listar("77770")], ["Zélia Manual"])
        self.assertEqual([s.nome for s in self.s.servidores.listar("111.222.333")], ["Zélia Manual"])
        self.assertEqual([s.nome for s in self.s.servidores.listar("11122233344")], ["Zélia Manual"])  # só dígitos
        self.s.servidores.remover(novo.id)
        self.assertEqual(self.s.servidores.listar("77770"), [])
        self.assertEqual(self.s.servidores.contar(), 3)
        self.assertEqual(self.s.servidores.obter(novo.id).nome, "Zélia Manual")  # histórico preservado
        # a mesma matrícula pode ser cadastrada de novo: reativa o cadastro
        self.s.servidores.salvar(Servidor(nome="Zélia Nova", matricula="77770"))
        self.assertEqual(self.s.servidores.listar("77770")[0].nome, "Zélia Nova")

    def test_matricula_duplicada_na_edicao(self):
        self.s.servidores.importar_planilha(self.planilha)
        ana, carlos = self.s.servidores.listar("ana")[0], self.s.servidores.listar("carlos")[0]
        carlos.matricula = ana.matricula
        with self.assertRaises(MatriculaDuplicada):
            self.s.servidores.salvar(carlos)

    def test_fluxo_gerar_e_enviar(self):
        r = self.s.gerar_folhas.executar(self.COMP, self.planilha)
        self.assertEqual(len(r.itens), 3)
        self.assertTrue(all(i.arquivo and i.arquivo_hash for i in r.itens))
        self.assertEqual(len(self.s.consultas.folhas_arquivadas("", self.COMP)), 3)

        selecionados = [i for i in r.itens if i.tem_email]  # Maria não tem e-mail
        resultado = self.s.envio.executar(r.op_id, self.COMP, TipoOperacao.GERACAO, selecionados, r.itens)
        self.assertEqual((resultado.enviados, resultado.total), (1, 2))
        self.assertIn("Carlos", resultado.erros[0])
        self.assertTrue(os.path.exists(resultado.relatorio))
        self.assertEqual(self.email.enviados, [("ana@teste.br", "Folha_17289104.pdf")])

        op = self.s.auditoria.listar()[0]
        self.assertEqual((op.status, op.qtd_enviados, op.qtd_erros, op.qtd_itens),
                         ("Concluída com erros de envio", 1, 1, 3))
        maria = next(i for i in self.s.auditoria.repositorio.itens(op.id) if i.nome == "Maria Silva")
        self.assertFalse(maria.selecionado)
        self.assertEqual(len(self.s.consultas.envios_realizados()), 2)

    def test_grupos_de_envio_e_reenvio_de_erros(self):
        r = self.s.gerar_folhas.executar(self.COMP, self.planilha)
        itens = {i.nome: i for i in r.itens}
        todos = self.s.envio.selecionar(r.itens, "todos", self.COMP)
        self.assertEqual({i.nome for i in todos}, {"Ana Beatriz Oliveira", "Carlos Eduardo Santos"})  # Maria sem e-mail
        self.s.envio.executar(r.op_id, self.COMP, TipoOperacao.GERACAO, todos, r.itens)
        # Ana enviada, Carlos com erro (endereço recusado)
        self.assertEqual([i.nome for i in self.s.envio.selecionar(r.itens, "erros", self.COMP)], ["Carlos Eduardo Santos"])
        self.assertEqual([i.nome for i in self.s.envio.selecionar(r.itens, "pendentes", self.COMP)], ["Carlos Eduardo Santos"])
        # histórico com busca por nome, matrícula e CPF
        self.assertEqual(len(self.s.consultas.envios_realizados("carlos")), 1)
        self.assertEqual(len(self.s.consultas.envios_realizados("17289104")), 1)
        self.assertEqual(len(self.s.consultas.envios_realizados("44455566604")), 1)
        self.assertEqual(len(self.s.consultas.envios_realizados("", "Erro")), 1)
        # o e-mail é corrigido e os erros são reenviados
        carlos = self.s.servidores.listar("carlos")[0]
        carlos.email = "carlos@certo.br"
        self.s.servidores.salvar(carlos)
        op_id, resultado = self.s.envio.reenviar_erros(self.s.consultas.envios_realizados("", "Erro"))
        self.assertEqual((resultado.enviados, resultado.total), (1, 1))
        self.assertIn(("carlos@certo.br", "Folha_17289105.pdf"), self.email.enviados)
        self.assertEqual(self.s.auditoria.repositorio.obter(op_id).tipo, TipoOperacao.REENVIO)
        with self.assertRaises(ErroDominio):  # nada mais com erro pendente
            self.s.envio.reenviar_erros(self.s.consultas.envios_realizados("", "Erro"))

    def test_fluxo_assinadas_com_conferencia(self):
        geracao = self.s.gerar_folhas.executar(self.COMP, self.planilha)
        pdfs = [i.arquivo for i in geracao.itens]
        sem_texto = os.path.join(self.pasta, "escaneada.pdf")
        doc = pymupdf.open()
        doc.new_page()
        doc.save(sem_texto)
        doc.close()

        op_id, folhas = self.s.folhas_assinadas.importar_lote(pdfs[:2] + [sem_texto])
        self.assertEqual([f.metodo for f in folhas][:2], [MetodoIdentificacao.TEXTO] * 2)
        anonima = folhas[2]
        self.assertIsNone(anonima.servidor_id)

        maria = self.s.servidores.listar("maria")[0]
        self.s.folhas_assinadas.atribuir(op_id, folhas, anonima, maria.id)
        self.assertEqual((anonima.servidor_id, anonima.metodo), (maria.id, MetodoIdentificacao.MANUAL))
        self.assertEqual(anonima.competencia, "")
        with self.assertRaises(ErroDominio):  # folha atribuída sem mês não pode ser arquivada
            self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)
        self.s.folhas_assinadas.definir_competencia(op_id, folhas, anonima, self.COMP)

        ana = self.s.servidores.listar("ana")[0]
        antiga = self.s.folhas_assinadas.folha_do_servidor(folhas, ana.id)
        nova = self.s.folhas_assinadas.substituir(op_id, folhas, ana.id, pdfs[2], self.COMP)
        self.assertIsNone(antiga.servidor_id)  # a folha anterior ficou sem dono
        self.assertEqual(nova.servidor_id, ana.id)

        itens = self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)
        self.assertEqual(len(itens), 3)
        assinadas = [f for f in self.s.consultas.folhas_arquivadas("", self.COMP) if f.assinada]
        self.assertEqual(len(assinadas), 3)

        relatorio = self.s.envio.concluir_sem_envio(op_id, self.s.folhas_assinadas.itens_auditoria(itens, folhas))
        self.assertTrue(os.path.exists(relatorio))
        eventos = [e.descricao for e in self.s.auditoria.repositorio.eventos(op_id)]
        self.assertTrue(any("atribuída manualmente" in e for e in eventos))
        self.assertTrue(any("substituída manualmente" in e for e in eventos))

    def test_reimportar_mesmo_arquivo_reaproveita_registro(self):
        geracao = self.s.gerar_folhas.executar(self.COMP, self.planilha)
        _, primeira = self.s.folhas_assinadas.importar_lote([geracao.itens[0].arquivo])
        _, segunda = self.s.folhas_assinadas.importar_lote([geracao.itens[0].arquivo])
        self.assertTrue(segunda[0].duplicado)
        self.assertEqual(primeira[0].doc_id, segunda[0].doc_id)

    def _pdf_folhas(self, nome, paginas):
        """PDF de várias páginas, cada uma com (servidor, matrícula, mês) — como um lote digitalizado."""
        caminho = os.path.join(self.pasta, nome)
        doc = pymupdf.open()
        for servidor, matricula, mes in paginas:
            doc.new_page().insert_text((50, 72), f"NOME DO SERVIDOR: {servidor}\nMATRICULA: {matricula}\n"
                                                 f"REFERENCIA: {mes}")
        doc.save(caminho)
        doc.close()
        return caminho

    def test_pdf_com_varias_folhas_e_meses_diferentes(self):
        self.s.servidores.importar_planilha(self.planilha)
        pdf = self._pdf_folhas("lote.pdf", [("Ana Beatriz Oliveira", "17289104", "SETEMBRO/2026"),
                                            ("Carlos Eduardo Santos", "17289105", "AGOSTO/2026"),
                                            ("Ana Beatriz Oliveira", "17289104", "AGOSTO/2026")])
        op_id, folhas = self.s.folhas_assinadas.importar_lote([pdf])
        self.assertEqual([(f.pagina, f.competencia) for f in folhas],
                         [(1, "SETEMBRO/2026"), (2, "AGOSTO/2026"), (3, "AGOSTO/2026")])
        self.assertTrue(all(f.servidor_id for f in folhas))
        self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)
        raiz = self.s.armazenamento.pasta_digitalizadas_raiz
        self.assertEqual(sorted(os.listdir(os.path.join(raiz, "Folhas_2026-08"))),
                         ["Ana_Beatriz_Oliveira-2026-08.pdf", "Carlos_Eduardo_Santos-2026-08.pdf"])
        self.assertEqual(os.listdir(os.path.join(raiz, "Folhas_2026-09")), ["Ana_Beatriz_Oliveira-2026-09.pdf"])
        self.assertEqual(self.s.auditoria.listar()[0].competencia, "VÁRIOS MESES")

    def test_pdf_ilegivel_avisa_o_usuario(self):
        self.s.servidores.importar_planilha(self.planilha)
        corrompido = os.path.join(self.pasta, "corrompido.pdf")
        with open(corrompido, "wb") as f:
            f.write(b"%PDF-1.4\n lixo")
        with self.assertRaises(ErroDominio) as e:
            self.s.folhas_assinadas.importar_lote([corrompido])
        self.assertIn("corrompido.pdf", str(e.exception))
        self.assertEqual(self.s.auditoria.listar()[0].status, "Falha na leitura")

    def test_servidores_sem_folha_considera_lotes_anteriores(self):
        self.s.servidores.importar_planilha(self.planilha)
        fa = self.s.folhas_assinadas
        op_id, folhas = fa.importar_lote([self._pdf_folhas("lote1.pdf", [("Ana Beatriz Oliveira", "17289104", "SETEMBRO/2026")])])
        fa.concluir_conferencia(op_id, folhas)
        _, folhas = fa.importar_lote([self._pdf_folhas("lote2.pdf", [("Carlos Eduardo Santos", "17289105", "SETEMBRO/2026")])])
        self.assertEqual([s.nome for s in fa.servidores_sem_folha(folhas, "SETEMBRO/2026")], ["Maria Silva"])
        self.assertEqual(len(fa.servidores_sem_folha(folhas, "AGOSTO/2026")), 3)

    def test_ocr_tenta_outras_variacoes_ate_achar_servidor_e_mes(self):
        leituras = ["", "NOME DO SERVIDOR: Ana Beatriz Oliveira", "REFERÊNCIA: SETEMBRO/2026", "nunca lida"]
        lidas = []

        class LeitorEscaneado(LeitorFolhasTesseract):
            texto_pagina = staticmethod(lambda caminho, pagina: "")
            ocr_disponivel = staticmethod(lambda: True)

            def leituras_ocr(self, caminho, pagina):
                for texto in leituras:
                    lidas.append(texto)
                    yield texto

        s = montar_servicos(self.pasta, servico_email=self.email, leitor_folhas=LeitorEscaneado())
        s.servidores.importar_planilha(self.planilha)
        pdf = os.path.join(self.pasta, "escaneada.pdf")
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(pdf)
        _, (folha,) = s.folhas_assinadas.importar_lote([pdf])
        self.assertEqual(s.servidores.obter(folha.servidor_id).nome, "Ana Beatriz Oliveira")
        self.assertEqual((folha.competencia, folha.metodo), ("SETEMBRO/2026", MetodoIdentificacao.OCR))
        self.assertEqual(len(lidas), 3)  # parou ao achar servidor e mês

    def test_foto_torta_tem_a_inclinacao_detectada(self):
        pdf = self._pdf_folhas("reta.pdf", [("Ana Beatriz Oliveira", "17289104", self.COMP)])
        with pymupdf.open(pdf) as doc:
            pix = doc[0].get_pixmap(dpi=150)
        folha = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        for angulo in (-12, -4, 7):
            torta = folha.rotate(angulo, expand=True, resample=Image.BICUBIC, fillcolor=(125, 98, 72))
            self.assertAlmostEqual(inclinacao(torta), -angulo, delta=1)

    def test_duplicidade_no_mesmo_pdf_vai_para_conferencia(self):
        self.s.servidores.importar_planilha(self.planilha)
        pdf = self._pdf_folhas("dup.pdf", [("Ana Beatriz Oliveira", "17289104", self.COMP),
                                           ("Ana Beatriz Oliveira", "17289104", self.COMP)])
        # conteúdos idênticos geram o mesmo hash: força páginas diferentes
        doc = pymupdf.open(pdf)
        doc[1].insert_text((50, 200), "segunda via")
        doc.saveIncr()
        doc.close()
        op_id, folhas = self.s.folhas_assinadas.importar_lote([pdf])
        self.assertEqual(len(folhas), 2)
        self.assertIsNotNone(folhas[0].servidor_id)
        self.assertIsNone(folhas[1].servidor_id)
        self.assertIn("Duplicada", folhas[1].observacao)
        # o usuário prefere a segunda via
        ana = self.s.servidores.listar("ana")[0]
        self.s.folhas_assinadas.atribuir(op_id, folhas, folhas[1], ana.id)
        self.assertIsNone(folhas[0].servidor_id)
        self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)

    def test_nova_versao_preserva_backup_old(self):
        self.s.servidores.importar_planilha(self.planilha)
        v1 = self._pdf_folhas("v1.pdf", [("Ana Beatriz Oliveira", "17289104", self.COMP)])
        op_id, folhas = self.s.folhas_assinadas.importar_lote([v1])
        self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)
        v2 = self._pdf_folhas("v2.pdf", [("Ana Beatriz Oliveira", "17289104", self.COMP)])
        doc = pymupdf.open(v2)
        doc[0].insert_text((50, 200), "versão revisada")
        doc.saveIncr()
        doc.close()
        op_id, folhas = self.s.folhas_assinadas.importar_lote([v2])
        self.assertIn("Substitui", folhas[0].observacao)
        self.s.folhas_assinadas.concluir_conferencia(op_id, folhas)
        pasta = os.path.join(self.s.armazenamento.pasta_digitalizadas_raiz, "Folhas_2026-02")
        arquivos = sorted(os.listdir(pasta))
        self.assertEqual(len(arquivos), 2)
        self.assertIn("Ana_Beatriz_Oliveira-2026-02.pdf", arquivos)
        antigo = next(a for a in arquivos if "_OLD_" in a)
        self.assertRegex(antigo, r"^Ana_Beatriz_Oliveira-2026-02_OLD_\d{8}-\d{6}\.pdf$")
        vigentes = [f for f in self.s.consultas.folhas_arquivadas("", self.COMP) if f.assinada]
        self.assertEqual(len(vigentes), 1)

    def test_aceita_somente_pdf(self):
        self.s.servidores.importar_planilha(self.planilha)
        imagem = os.path.join(self.pasta, "foto.jpg")
        with open(imagem, "wb") as f:
            f.write(b"\xff\xd8\xff")
        falso_pdf = os.path.join(self.pasta, "falso.pdf")
        with open(falso_pdf, "wb") as f:
            f.write(b"\xff\xd8\xff")
        for caminho in (imagem, falso_pdf):
            with self.assertRaises(ErroDominio):
                self.s.folhas_assinadas.importar_lote([caminho])

    def _planilha_calendario(self, nome, linhas, cabecalho=("Data", "Especificação")):
        caminho = os.path.join(self.pasta, nome)
        wb = openpyxl.Workbook()
        wb.active.append(list(cabecalho))
        for linha in linhas:
            wb.active.append(list(linha))
        wb.save(caminho)
        return caminho

    def test_calendario_importacao_com_resumo(self):
        from datetime import date
        plan = self._planilha_calendario("cal.xlsx", [
            ("01/01/2027", "Feriado Nacional"), (date(2027, 7, 1), "Recesso Acadêmico"),
            ("31/02/2027", "Data impossível"), ("01/01/2027", "Repetida"), ("02/01/2027", ""),
        ])
        resumo = self.s.calendario.analisar_planilha(plan)
        self.assertEqual((resumo.novos, resumo.atualizados, len(resumo.ignoradas)), (2, 0, 3))
        self.assertEqual(self.s.calendario.listar(), [])  # nada gravado antes de confirmar
        self.s.calendario.confirmar_importacao(resumo)
        self.assertEqual([(d.data_br, d.especificacao) for d in self.s.calendario.listar(2027)],
                         [("01/01/2027", "Feriado Nacional"), ("01/07/2027", "Recesso Acadêmico")])
        nova = self._planilha_calendario("cal2.xlsx", [("01/01/2027", "Confraternização Universal"),
                                                       ("01/07/2027", "Recesso Acadêmico")])
        r2 = self.s.calendario.analisar_planilha(nova)
        self.assertEqual((r2.novos, r2.atualizados, r2.sem_alteracao), (0, 1, 1))

    def test_calendario_cabecalho_e_formato(self):
        ruim = self._planilha_calendario("ruim.xlsx", [("01/01/2027", "x")], cabecalho=("Dia", "Nota"))
        with self.assertRaises(PlanilhaInvalida):
            self.s.calendario.analisar_planilha(ruim)
        csv = os.path.join(self.pasta, "cal.csv")
        with open(csv, "w") as f:
            f.write("Data;Especificação")
        with self.assertRaises(PlanilhaInvalida):
            self.s.calendario.analisar_planilha(csv)

    def test_calendario_manual(self):
        from datetime import date
        self.s.calendario.salvar(date(2026, 2, 16), "Ponto Facultativo")
        self.s.calendario.salvar(date(2026, 2, 16), "Carnaval")  # altera a mesma data
        self.assertEqual([d.especificacao for d in self.s.calendario.listar(2026)], ["Carnaval"])
        self.assertEqual(self.s.calendario.do_mes("FEVEREIRO/2026"), {16: "Carnaval"})
        self.assertEqual(self.s.calendario.do_mes("MARÇO/2026"), {})
        with self.assertRaises(ErroDominio):
            self.s.calendario.salvar(date(2026, 2, 17), "   ")
        self.s.calendario.remover(date(2026, 2, 16))
        self.assertEqual(self.s.calendario.listar(), [])

    def test_configuracao_email(self):
        cfg = ConfiguracaoEmail(metodo="Outlook", host="h", porta="465", usuario="u", senha="p", usar_tls=False)
        self.s.configuracoes.salvar_email(cfg)
        self.assertEqual(self.s.configuracoes.email(), cfg)
        self.s.configuracoes.repositorio.definir("metodo_envio", "Resend")  # método antigo, removido
        self.assertEqual(self.s.configuracoes.email().metodo, "SMTP")


if __name__ == "__main__":
    unittest.main()
