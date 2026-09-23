"""Preenchimento do modelo institucional real (.docx)."""
import os
import tempfile
import unittest

import docx

from folha_ponto.dominio.entidades import Servidor
from folha_ponto.infraestrutura.folha_docx import _celulas_unicas, preencher_folha

MODELO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      "Modelo de folha de ponto - Exemplo.docx")


class TestFolhaDocx(unittest.TestCase):
    def preencher(self, competencia, **dados):
        destino = os.path.join(tempfile.mkdtemp(), "folha.docx")
        servidor = Servidor(nome="Patrícia Alves", matricula="17289110", carga_horaria="40h", **dados)
        preencher_folha(MODELO, servidor, competencia, destino)
        return docx.Document(destino).tables[0]

    def test_cabecalho(self):
        tabela = self.preencher("FEVEREIRO/2026", cargo="")
        texto = "\n".join(c.text for row in tabela.rows[:3] for c in _celulas_unicas(row))
        self.assertIn("REFERÊNCIA: FEVEREIRO/2026", texto)
        self.assertIn("NOME DO SERVIDOR: PATRÍCIA ALVES", texto)
        self.assertIn("MATRÍCULA: 17289110", texto)
        self.assertIn("PROF-M20", texto)  # cargo vazio usa o padrão do modelo
        rotulo = next(r for c in _celulas_unicas(tabela.rows[2]) for p in c.paragraphs for r in p.runs
                      if r.text == "NOME DO SERVIDOR")
        self.assertTrue(rotulo.bold)

    def test_calendario_do_mes(self):
        tabela = self.preencher("FEVEREIRO/2026")
        dias = {int(c[0].text): c[1].text.strip() for c in (_celulas_unicas(r) for r in tabela.rows)
                if len(c) == 7 and c[0].text.strip().isdigit()}
        self.assertEqual(dias[1], "DOMINGO")   # 01/02/2026 é domingo
        self.assertEqual(dias[7], "SÁBADO")
        self.assertEqual(dias[2], "")
        self.assertEqual(max(dias), 28)          # 29, 30 e 31 ficam em branco


if __name__ == "__main__":
    unittest.main()
