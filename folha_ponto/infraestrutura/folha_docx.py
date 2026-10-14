"""Preenchimento do modelo institucional da folha de ponto (.docx) e geração em lote."""
import calendar
import os

import docx

from ..dominio.calendario import rotulo_na_folha
from ..dominio.competencia import separar_competencia
from ..dominio.entidades import Servidor
from ..dominio.regras import nome_arquivo_seguro
from .conversor_pdf import ConversorPDF


def _celulas_unicas(row):
    """Células de uma linha sem repetição (células mescladas aparecem várias vezes em row.cells)."""
    vistas, unicas = set(), []
    for c in row.cells:
        if id(c._tc) not in vistas:
            vistas.add(id(c._tc))
            unicas.append(c)
    return unicas


def _escrever(paragrafo, partes):
    """Reescreve o parágrafo preservando a fonte do 1º trecho. partes = [(texto, negrito), ...]."""
    tamanho = paragrafo.runs[0].font.size if paragrafo.runs else None
    nome_fonte = paragrafo.runs[0].font.name if paragrafo.runs else None
    for r in list(paragrafo.runs):
        r._r.getparent().remove(r._r)
    for texto, negrito in partes:
        run = paragrafo.add_run(texto)
        run.bold = negrito
        if tamanho:
            run.font.size = tamanho
        if nome_fonte:
            run.font.name = nome_fonte


def _texto_celula(celula, texto, negrito=True):
    for extra in celula.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _escrever(celula.paragraphs[0], [(texto, negrito)])


def _preencher_calendario(tabela, competencia, especiais=None):
    """Reescreve a grade de dias ("01".."31") conforme o mês/ano: marca sábados/domingos e datas especiais
    (feriado, recesso, ponto facultativo) e deixa em branco os dias que não existem no mês (ex: 31/04, 30/02)."""
    especiais = especiais or {}
    mes, ano = separar_competencia(competencia)
    if not mes:
        return
    dias_no_mes = calendar.monthrange(ano, mes)[1]
    for row in tabela.rows:
        celulas = _celulas_unicas(row)
        if len(celulas) != 7:
            continue
        rotulo = celulas[0].text.strip()
        if not (rotulo.isdigit() and 1 <= int(rotulo) <= 31):
            continue
        dia = int(rotulo)
        if dia > dias_no_mes:
            valores = [""] * 7
        else:
            semana = calendar.weekday(ano, mes, dia)  # 5 = sábado, 6 = domingo
            if semana == 5:
                valores = [f"{dia:02d} ", "SÁBADO ", "-- ", "-- ", "SÁBADO ", "-- ", "-- "]
            elif semana == 6:
                valores = [f"{dia:02d} ", "DOMINGO ", "-- ", "-- ", "DOMINGO ", "-- ", "-- "]
            elif dia in especiais:
                rotulo_dia = rotulo_na_folha(especiais[dia])
                valores = [f"{dia:02d} ", f"{rotulo_dia} ", "-- ", "-- ", f"{rotulo_dia} ", "-- ", "-- "]
            else:
                valores = [f"{dia:02d} ", " ", " ", " ", " ", "", ""]
        for celula, valor in zip(celulas, valores):
            _texto_celula(celula, valor)


def _preencher_observacoes(tabela, competencia, especiais):
    """Escreve na linha seguinte a “Observações:” as datas especiais do mês (ex.: "Dia 7, Feriado Nacional")."""
    if not especiais:
        return
    mes, _ = separar_competencia(competencia)
    notas = "; ".join(f"Dia {dia}, {texto}" for dia, texto in sorted(especiais.items()))
    linhas = list(tabela.rows)
    for i, row in enumerate(linhas):
        if "observa" in "".join(c.text for c in _celulas_unicas(row)).lower() and i + 1 < len(linhas):
            celula = _celulas_unicas(linhas[i + 1])[0]
            _texto_celula(celula, notas, negrito=False)
            return


def _campo(servidor, nome):
    """Valor do servidor ou, se vazio, o padrão do modelo institucional (ex.: cargo de professor)."""
    return (getattr(servidor, nome) or Servidor.__dataclass_fields__[nome].default or "").strip()


def preencher_folha(modelo, servidor, competencia, destino, datas_especiais=None):
    """Preenche o cabeçalho (mantendo os rótulos em negrito) e o calendário do mês.
    datas_especiais: {dia: especificação} dos feriados/recessos do mês."""
    documento = docx.Document(modelo)
    ua, cargo, padrao = _campo(servidor, "ua"), _campo(servidor, "cargo"), _campo(servidor, "padrao")
    funcao, exercicio = _campo(servidor, "funcao"), _campo(servidor, "exercicio")
    for tabela in documento.tables:
        for row in tabela.rows:
            for celula in _celulas_unicas(row):
                for p in celula.paragraphs:
                    texto = p.text
                    if "REFERÊNCIA:" in texto or "REFERENCIA:" in texto:
                        _escrever(p, [(f"REGISTRO DE FREQUÊNCIA           REFERÊNCIA: {competencia.upper()} ", True)])
                    elif "MATRÍCULA:" in texto or "MATRICULA:" in texto:
                        _escrever(p, [("UA", True), (f":  {ua}                     ", False),
                                      ("MATRÍCULA: ", True), (f"{servidor.matricula} ", False)])
                    elif "NOME DO SERVIDOR" in texto:
                        _escrever(p, [("NOME DO SERVIDOR", True), (f": {servidor.nome.upper()} ", False)])
                    elif "CARGO:" in texto:
                        _escrever(p, [("CARGO:", True), (f"  {cargo}              ", False), ("PADRÃO:", True),
                                      (f" {padrao}\n", False), ("FUNÇÃO", True),
                                      (f": {funcao}                                                              ", False),
                                      ("CARGA HORÁRIA:", True), (f"  {servidor.carga_horaria}", False)])
                    elif "EXERCÍCIO" in texto or "EXERCICIO" in texto:
                        _escrever(p, [("EXERCÍCIO", True), (f": {exercicio} ", False)])
        _preencher_calendario(tabela, competencia, datas_especiais)
        _preencher_observacoes(tabela, competencia, datas_especiais)
    os.makedirs(os.path.dirname(os.path.abspath(destino)), exist_ok=True)
    documento.save(destino)
    return destino


class GeradorFolhasDocx:
    def __init__(self, armazenamento):
        self.armazenamento = armazenamento

    def gerar_lote(self, servidores, competencia, progresso, datas_especiais=None):
        """-> ({servidor_id: (docx, pdf ou None)}, erros)"""
        pasta = self.armazenamento.pasta_geradas(competencia)
        arquivos, erros = {}, []
        with ConversorPDF() as conversor:
            for n, s in enumerate(servidores):
                progresso(n, len(servidores), f"Gerando folha de {s.nome}")
                base = os.path.join(pasta, f"Folha_{nome_arquivo_seguro(s.matricula)}_{nome_arquivo_seguro(s.nome)}")
                try:
                    preencher_folha(self.armazenamento.modelo_folha, s, competencia, base + ".docx", datas_especiais)
                    pdf_ok = conversor.converter(base + ".docx", base + ".pdf")
                    arquivos[s.id] = (base + ".docx", base + ".pdf" if pdf_ok else None)
                except Exception as e:
                    erros.append(f"{s.nome}: {e}")
        return arquivos, erros
